"""Generate the 23-page V0.12 academic defence deck from one design result.

All tables, charts and simple bridge diagrams are native PowerPoint objects.
Every chart therefore carries an editable embedded workbook.  The numerical
content is read only from the calculation result used by the report writer.
"""
from math import isfinite
import re
import tempfile
import zipfile
from pathlib import Path
from io import BytesIO

import numpy as np
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from lxml import etree

from .geometry import Geometry
from .units import bridge_description

INDIGO = "1F3864"
RED = "9E2A2B"
DARK = "595959"
MID = "9DA9BC"
LIGHT = "EEF1F5"
WHITE = "FFFFFF"
SLIDE_COUNT = 23


def rgb(value):
    return RGBColor.from_string(value)


def _font_run(run, size=13, bold=False, color=DARK, east="SimSun", latin="Times New Roman"):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    for tag, face in (("a:latin", latin), ("a:ea", east), ("a:cs", latin)):
        node = rpr.find(qn(tag))
        if node is None:
            node = OxmlElement(tag); rpr.append(node)
        node.set("typeface", face)


def _style_text_frame(tf, size=13, bold=False, color=DARK, east="SimSun", latin="Times New Roman",
                      align=PP_ALIGN.LEFT):
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for paragraph in tf.paragraphs:
        paragraph.alignment = align
        paragraph.space_after = Pt(3)
        paragraph.line_spacing = 1.05
        for run in paragraph.runs:
            _font_run(run, size, bold, color, east, latin)


def _textbox(slide, x, y, w, h, text, size=13, color=DARK, bold=False,
             east="SimSun", latin="Times New Roman", align=PP_ALIGN.LEFT,
             valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.clear()
    box.text_frame.vertical_anchor = valign
    for index, line in enumerate(str(text).split("\n")):
        p = box.text_frame.paragraphs[0] if index == 0 else box.text_frame.add_paragraph()
        p.text = line
    _style_text_frame(box.text_frame, size, bold, color, east, latin, align)
    return box


def _line(slide, x1, y1, x2, y2, color=INDIGO, width=1.25, dash=None):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                       Inches(x2), Inches(y2))
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    if dash is not None:
        shape.line.dash_style = dash
    return shape


def _rect(slide, x, y, w, h, fill=WHITE, line_color=INDIGO, line_width=1.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line_color); shape.line.width = Pt(line_width)
    return shape


def _set_cell_border(cell, top=None, bottom=None):
    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB", "a:lnTlToBr", "a:lnBlToTr"):
        old = tc_pr.find(qn(tag))
        if old is not None:
            tc_pr.remove(old)
    for tag, spec in (("a:lnT", top), ("a:lnB", bottom)):
        if spec is None:
            continue
        ln = OxmlElement(tag); ln.set("w", str(int(spec[1] * 12700))); tc_pr.append(ln)
        solid = OxmlElement("a:solidFill"); ln.append(solid)
        srgb = OxmlElement("a:srgbClr"); solid.append(srgb)
        srgb.set("val", spec[0])
        dash = OxmlElement("a:prstDash"); dash.set("val", "solid"); ln.append(dash)


def _table(slide, caption, heads, rows, x=0.62, y=1.72, w=12.09, h=4.5,
           widths=None, font_size=11):
    _textbox(slide, x, y - 0.34, w, 0.28, caption, 10.5, INDIGO, True, "SimSun")
    table = slide.shapes.add_table(len(rows) + 1, len(heads), Inches(x), Inches(y), Inches(w), Inches(h)).table
    if widths:
        for col, width in zip(table.columns, widths):
            col.width = Inches(width)
    for i, row in enumerate([heads] + rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(value)
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(WHITE)
            cell.margin_left = cell.margin_right = Inches(0.05)
            cell.margin_top = cell.margin_bottom = Inches(0.025)
            _style_text_frame(cell.text_frame, font_size, i == 0, INDIGO if i == 0 else DARK,
                              "SimHei" if i == 0 else "SimSun", "Times New Roman",
                              PP_ALIGN.CENTER if j else PP_ALIGN.LEFT)
            top = (INDIGO, 1.25) if i == 0 else None
            bottom = (INDIGO, 0.8) if i == 0 else ((INDIGO, 1.25) if i == len(rows) else None)
            _set_cell_border(cell, top, bottom)
    return table


def _format_chart(chart, y_title, x_title="纵向坐标 x / m", legend=True):
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.name = "SimSun"; chart.legend.font.size = Pt(10.5)
    chart.has_title = False
    chart.category_axis.has_title = bool(x_title)
    chart.value_axis.has_title = bool(y_title)
    if x_title:
        chart.category_axis.axis_title.text_frame.text = x_title
    if y_title:
        chart.value_axis.axis_title.text_frame.text = y_title
    for axis in (chart.category_axis, chart.value_axis):
        axis.tick_labels.font.name = "Times New Roman"
        axis.tick_labels.font.size = Pt(10.5)
        if axis.has_title:
            _style_text_frame(axis.axis_title.text_frame, 10.5, False, DARK, "SimSun")
        axis.format.line.color.rgb = rgb(MID)
        axis.format.line.width = Pt(0.75)
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = rgb(LIGHT)
    return chart


def _xy_chart(slide, caption, series, y_title, x=0.72, y=1.88, w=11.88, h=4.72,
              x_title="纵向坐标 x / m"):
    data = XyChartData()
    for name, xs, ys in series:
        s = data.add_series(name)
        for xv, yv in zip(xs, ys):
            if isfinite(float(xv)) and isfinite(float(yv)):
                s.add_data_point(round(float(xv), 8), round(float(yv), 8))
    chart = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS,
                                   Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    _format_chart(chart, y_title, x_title)
    for idx, s in enumerate(chart.series):
        s.format.line.color.rgb = rgb([INDIGO, RED, MID][idx % 3])
        s.format.line.width = Pt(1.7)
    _textbox(slide, x, y + h + 0.08, w, 0.25, caption, 10.5, DARK, False, "SimSun", align=PP_ALIGN.CENTER)
    return chart


def _bar_chart(slide, caption, categories, series, y_title, x=0.75, y=1.92, w=11.83, h=4.62):
    data = CategoryChartData(); data.categories = [str(v) for v in categories]
    for name, values in series:
        data.add_series(name, [round(float(v), 8) for v in values])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w),
                                   Inches(h), data).chart
    _format_chart(chart, y_title, "")
    chart.category_axis.tick_labels.font.name = "SimSun"
    for idx, s in enumerate(chart.series):
        s.format.fill.solid(); s.format.fill.fore_color.rgb = rgb([INDIGO, RED, MID][idx % 3])
        s.format.line.color.rgb = rgb([INDIGO, RED, MID][idx % 3])
    _textbox(slide, x, y + h + 0.08, w, 0.25, caption, 10.5, DARK, False, "SimSun", align=PP_ALIGN.CENTER)
    return chart


def _add_notes(slide, source, focus):
    slide.notes_slide.notes_text_frame.text = f"数据来源：{source}\n讲解重点：{focus}"


def _control_rows(controls):
    wanted = {"左边跨跨中", "左主墩", "主跨四分点", "主跨跨中", "右主墩"}
    return [v for v in controls if v["name"] in wanted]


def _finalize_chart_package(path):
    """Normalize axis ids and synchronize chart caches with embedded workbooks."""
    source = Path(path)
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pptx", dir=source.parent) as handle:
        temporary = Path(handle.name)
    with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            content = zin.read(item.filename)
            if item.filename.startswith("ppt/charts/chart") and item.filename.endswith(".xml"):
                chart_number = re.search(r"chart(\d+)\.xml$", item.filename).group(1)
                workbook_name = f"ppt/embeddings/Microsoft_Excel_Sheet{chart_number}.xlsx"
                values = {}
                if workbook_name in zin.namelist():
                    with zipfile.ZipFile(BytesIO(zin.read(workbook_name)), "r") as workbook:
                        sheet = etree.fromstring(workbook.read("xl/worksheets/sheet1.xml"))
                        ns_sheet = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                        for cell in sheet.xpath(".//s:c[s:v]", namespaces=ns_sheet):
                            values[cell.get("r")] = cell.find("s:v", namespaces=ns_sheet).text
                root = etree.fromstring(content)
                ns_chart = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
                for axis in root.xpath(".//c:axId | .//c:crossAx", namespaces=ns_chart):
                    value = int(axis.get("val"))
                    if value < 0:
                        axis.set("val", str(4294967296 + value))
                for num_ref in root.xpath(".//c:numRef", namespaces=ns_chart):
                    formula = num_ref.find("c:f", namespaces=ns_chart)
                    cache = num_ref.find("c:numCache", namespaces=ns_chart)
                    if formula is None or cache is None:
                        continue
                    match = re.search(r"\$([A-Z]+)\$(\d+):\$[A-Z]+\$(\d+)$", formula.text or "")
                    if not match:
                        continue
                    column, start, end = match.group(1), int(match.group(2)), int(match.group(3))
                    for point in cache.findall("c:pt", namespaces=ns_chart):
                        address = f"{column}{start + int(point.get('idx'))}"
                        value_node = point.find("c:v", namespaces=ns_chart)
                        if value_node is not None and address in values:
                            value_node.text = values[address]
                content = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, content)
    temporary.replace(source)


def export_slides(result, path):
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    c = result["input"]; g = result["geometry"]; a = result["analysis"]
    loads = result["loads"]; p = result["prestress"]; cs = result["construction_stages"]
    sv = result["structural_verification"]; ck = sv["checks"]; summary = result["presentation_summary"]
    geometry = Geometry(c)
    desc = bridge_description(c, g)
    lanes = desc["lanes_per_deck"]
    deck_model = "两幅独立纵向平面梁" if desc["deck_count"] == 2 else "单幅双向纵向平面梁"

    def new(title, conclusion, source, focus):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _textbox(slide, 0.62, 0.30, 12.09, 0.45, title, 24, INDIGO, True, "SimHei")
        _line(slide, 0.62, 0.88, 12.71, 0.88, INDIGO, 1.25)
        _textbox(slide, 0.62, 1.02, 12.09, 0.42, conclusion, 13, RED, True, "SimHei")
        page = len(prs.slides)
        _line(slide, 0.62, 6.96, 12.71, 6.96, MID, 0.6)
        _textbox(slide, 0.62, 7.02, 10.8, 0.21, c["project_name"], 9, DARK, False, "SimSun")
        _textbox(slide, 11.55, 7.02, 1.16, 0.21, f"{page}/{SLIDE_COUNT}", 9, DARK, False,
                 "SimSun", align=PP_ALIGN.RIGHT)
        _add_notes(slide, source, focus)
        return slide

    slide = new(desc["title"] + "课程设计",
                f"本方案按 {g['left_side_span']:g}+{g['spans'][1]:g}+{g['right_side_span']:g} m {deck_model}进行课程范围验算。",
                "计算书封面、第1章；设计结果JSON", "说明方案对象、课程范围和汇报结构。")
    _textbox(slide, 0.75, 2.0, 11.8, 0.78, " + ".join(f"{value:g}" for value in g["spans"]) + " m", 36, INDIGO, True, "SimHei", align=PP_ALIGN.CENTER)
    _textbox(slide, 1.4, 3.05, 10.5, 0.85, f"{desc['arrangement']}\n单幅宽 {g['width']:.2f} m｜总宽 {g['total_width']:.2f} m", 18, DARK, False, "SimSun", align=PP_ALIGN.CENTER)
    _textbox(slide, 2.2, 4.42, 8.9, 0.56, f"课程设计状态：{result['status_code']}｜具名验算 {sv['summary']['passed']}/{sv['summary']['total']} 通过", 16, RED, True, "SimHei", align=PP_ALIGN.CENTER)
    _textbox(slide, 2.0, 5.42, 9.3, 0.48, "BridgeGenerator V0.12｜计算书、DWG 与 PPT 同源生成", 12, DARK, False, "SimSun", align=PP_ALIGN.CENTER)

    slide = new("设计任务与计算边界", "成果面向课程设计，结论仅覆盖纵向平面梁与列明的六项验算。", "计算书第1章、9.5节", "先说明任务边界，再解释OUT_OF_COURSE_SCOPE不等同于计算失败。")
    _table(slide, "表1-1 设计任务与边界", ["类别", "本版内容"], [["已计算", "恒载、汽车荷载、组合包络、预应力、双主墩施工阶段、承载能力与正常使用"], ["结构对象", desc['arrangement'] + "；纵向平面梁模型"], ["未列数值", "横向板壳、空间扭转、局部锚固区、风、地震、下部结构与基础承载力"], ["结果含义", f"{result['status_code']}：具名验算 {sv['summary']['passed']}/{sv['summary']['total']} 通过"]], h=3.65, widths=[1.8, 10.29], font_size=12)
    _textbox(slide, 0.75, 5.72, 11.7, 0.72, "▍边界  范围外项目保留 OUT_OF_COURSE_SCOPE 标记，不补造工程复核数据。", 13, DARK, False, "SimSun")

    slide = new("设计依据与计算模型", "全桥作用效应采用变 EI 连续梁有限元，荷载、组合和验算使用同一份离散模型。", "计算书第1章，表1-1、表1-2", "讲清模型自由度、材料、规范来源与重要性系数。")
    _table(slide, "表1-2 主要依据与模型参数", ["项目", "采用值", "来源"], [["结构模型", "Euler-Bernoulli 变 EI 连续梁", "程序计算模型"], ["混凝土", c["concrete"], "JTG 3362—2018"], ["汽车荷载", loads["traffic"]["class_name"], "JTG D60—2015"], ["结构重要性系数", f"{c['importance_factor']:.2f}", "基本组合"], ["纵向网格", f"最大 {c['mesh_length']:.2f} m，并强制包含跨中及四分点", "程序离散规则"]], h=3.9, widths=[2.25, 4.1, 5.74], font_size=11.5)
    _textbox(slide, 0.75, 5.94, 11.7, 0.52, "▍一致性  截面多边形同时用于面积、惯性矩、自重和 DWG 轮廓。", 13, DARK, False, "SimSun")

    slide = new("跨径、桥宽与结构布置", "自动边跨按主跨比例匹配；" + desc['arrangement'] + "。", "计算书第1章、第2章；总体布置图", "指出自动边跨逻辑和实际横断面布置。")
    x0, y0, width = 0.85, 3.08, 11.55; total = sum(g["spans"]); support_x = [x0 + width * value / total for value in g["supports"]]
    _line(slide, x0, y0, x0 + width, y0, INDIGO, 2.0)
    for idx, sx in enumerate(support_x):
        _line(slide, sx, y0, sx, 4.08, DARK, 1.2); _rect(slide, sx - 0.15, 4.08, 0.30, 0.18, LIGHT, DARK, 0.8)
        _textbox(slide, sx - 0.25, 4.38, 0.5, 0.25, "A" if idx in (0, 3) else f"P{idx}", 10.5, INDIGO, True, "SimHei", align=PP_ALIGN.CENTER)
    for idx, span in enumerate(g["spans"]):
        mid = (support_x[idx] + support_x[idx + 1]) / 2; _line(slide, support_x[idx], 2.55, support_x[idx + 1], 2.55, MID, 0.8)
        _textbox(slide, mid - 0.65, 2.26, 1.3, 0.25, f"{span:g} m", 11, DARK, False, "SimSun", align=PP_ALIGN.CENTER)
    _textbox(slide, 0.78, 1.82, 11.7, 0.42, f"▍布置  主跨 {g['spans'][1]:g} m；边跨比 {g['left_side_span']/g['spans'][1]:.3f} / {g['right_side_span']/g['spans'][1]:.3f}；总长 {g['total_length']:.1f} m。", 13, DARK, False, "SimSun")
    _textbox(slide, 0.78, 5.1, 11.7, 0.85, f"▍横断面  单幅 {g['width']:.2f} m，幅数 {desc['deck_count']}，两幅净距 {desc['gap']:.2f} m，总宽 {g['total_width']:.2f} m。" + ("左右幅镜像布置，计算结果为单幅代表值。" if desc['deck_count'] == 2 else "双向车道使用同一主梁，单幅不计两幅净距。"), 13, DARK, False, "SimSun")
    _textbox(slide, 0.78, 6.38, 11.7, 0.24, "图2-1 连续梁总体布置示意（PowerPoint 原生可编辑）", 10.5, DARK, False, "SimSun", align=PP_ALIGN.CENTER)

    slide = new("变截面箱梁几何设计", "梁高与底板、腹板厚度沿纵向联动，截面包含顶板加腋和底板倒角。", "计算书第2章；表2-1", "说明跨中与墩顶的变高关系及倒角参数。")
    outer, voids = geometry.outlines(g["supports"][1]); all_pts = outer + [q for poly in voids for q in poly]
    min_x = min(v[0] for v in all_pts); max_x = max(v[0] for v in all_pts); min_y = min(v[1] for v in all_pts); max_y = max(v[1] for v in all_pts); scale = min(6.2 / (max_x - min_x), 3.8 / (max_y - min_y))
    def draw_poly(points, color, width=1.5):
        mapped = [(0.9 + (xx - min_x) * scale, 2.0 + (max_y - yy) * scale) for xx, yy in points]
        for q1, q2 in zip(mapped, mapped[1:] + mapped[:1]): _line(slide, q1[0], q1[1], q2[0], q2[1], color, width)
    draw_poly(outer, INDIGO, 1.8)
    for poly in voids: draw_poly(poly, RED, 1.3)
    _textbox(slide, 7.75, 1.83, 4.65, 3.9, f"▍梁高  跨中 {g['height_mid']:.3f} m；墩顶 {g['height_pier']:.3f} m\n▍变高  指数 {g['curve_exponent']:.2f}；参考跨径 {c['height_reference_span']:.1f} m\n▍顶板  {c['top_thickness']*1000:.0f} mm；加腋 {c['top_haunch_horizontal']*1000:.0f}×{c['top_haunch_vertical']*1000:.0f} mm\n▍底板  {c['bottom_mid']*1000:.0f}～{c['bottom_pier']*1000:.0f} mm\n▍倒角  {c['bottom_chamfer_horizontal']*1000:.0f}×{c['bottom_chamfer_vertical']*1000:.0f} mm\n▍腹板  {c['web_mid']*1000:.0f}～{c['web_pier']*1000:.0f} mm", 13, DARK, False, "SimSun")
    _textbox(slide, 0.78, 6.36, 6.9, 0.25, "图2-2 墩顶单箱单室截面示意（含顶板加腋与底板倒角）", 10.5, DARK, False, "SimSun", align=PP_ALIGN.CENTER)

    controls = _control_rows(result["controls"])
    slide = new("截面特性与自重基础", "截面面积和惯性矩随梁高变化，墩顶刚度显著高于跨中。", "计算书第2章，表2-2", "说明所有刚度、自重和应力均来自同一多边形积分。")
    _table(slide, "表2-2 主要控制截面特性", ["截面", "x / m", "梁高 / m", "面积 / m²", "惯性矩 / m⁴"], [[v["name"], f"{v['x']:.2f}", f"{v['height']:.3f}", f"{v['area']:.3f}", f"{v['inertia']:.3f}"] for v in controls], h=4.25, widths=[3.0, 1.45, 2.0, 2.0, 3.64], font_size=11)
    _textbox(slide, 0.75, 6.18, 11.7, 0.45, "▍方法  外轮廓面积减去带倒角箱室面积，形心与惯性矩按多边形精确积分。", 13, DARK, False, "SimSun")

    d2 = loads["dead2_components"]
    slide = new("恒载组成与计算值", "一期恒载由变截面自重形成，二期恒载按铺装、护栏及附属恒载叠加。", "计算书第3章，表3-1、表3-2", "区分一期和二期恒载并指出单位。")
    _bar_chart(slide, "图3-1 单幅二期恒载组成", list(d2.keys()), [("线荷载", list(d2.values()))], "线荷载 / (kN·m⁻¹)", x=0.72, y=1.9, w=7.5, h=4.55)
    _textbox(slide, 8.55, 1.95, 3.8, 3.7, f"▍一期恒载  {loads['dead1_min']:.2f}～{loads['dead1_max']:.2f} kN/m\n▍铺装厚度  {c['pavement_thickness']*1000:.0f} mm\n▍混凝土重度  {c['concrete_weight']:.1f} kN/m³\n▍铺装重度  {c['pavement_weight']:.1f} kN/m³\n▍单幅总恒载  {loads['total_dead_weight_per_deck']:.1f} kN", 13, DARK, False, "SimSun")

    slide = new("恒载内力与支点反力", "恒载总重与支反力满足力平衡，内力用于各组合的永久作用项。", "计算书第3章，表3-3、表3-4；第5章，表5-7", "说明平衡误差和支点负弯矩特征。")
    _xy_chart(slide, "图3-2 单幅恒载弯矩", [("恒载弯矩", a["x"], np.asarray(a["M_dead"]) / 1000)], "弯矩 / (MN·m)", x=0.72, y=1.84, w=7.55, h=4.62)
    rx = result["reactions"]
    _table(slide, "表5-7 单幅恒载反力", ["支点", "一期 / kN", "二期 / kN", "合计 / kN"], [[rx["names"][i], f"{rx['dead1'][i]:.1f}", f"{rx['dead2'][i]:.1f}", f"{rx['dead'][i]:.1f}"] for i in range(len(rx["names"]))], x=8.48, y=2.0, w=4.2, h=3.35, widths=[1.05, 1.05, 1.05, 1.05], font_size=9.5)
    _textbox(slide, 8.55, 5.65, 3.95, 0.55, f"▍平衡误差  {result['verification']['force_balance_error_kn']:.3e} kN", 12, RED, True, "SimHei")

    traffic = loads["traffic"]
    full_case = max(traffic['transverse_cases'], key=lambda case: case['vertical_multiplier'])
    slide = new("公路Ⅰ级汽车荷载模型", f"车道荷载含均布荷载、集中荷载和冲击系数，单幅最多同时布置{lanes}车道。", "计算书第4章，表4-1；JTG D60—2015", "说明Pk随跨径取值、实际车道系数以及冲击系数。")
    _table(slide, "表4-1 单幅汽车荷载参数", ["参数", "采用值", "说明"], [["均布荷载 qk", f"{traffic['qk']:.3f} kN/m", "沿不利影响线区段布置"], ["集中荷载 Pk", f"{traffic['pk']:.3f} kN", "按主跨跨径计算"], ["基频", f"{traffic['frequency_hz']:.3f} Hz", "按成桥模型估计"], ["冲击系数 μ", f"{traffic['impact']:.4f}", "仅标准及基本组合计入"], ["全车道倍率 nξ", f"{lanes} × {full_case['lane_factor']:g} = {full_case['vertical_multiplier']:g}", "其他横向工况按实际启用车道数取系数"], ["横向计算", traffic["transverse_method"], "逐工况分别求包络"]], h=4.15, widths=[2.5, 3.1, 6.49], font_size=11.5)
    _textbox(slide, 0.75, 6.05, 11.7, 0.52, "▍纵向布载  正、负影响线区段独立取最不利加载，集中荷载在影响线极值处移动。", 13, DARK, False, "SimSun")

    slide = new("横向车道工况与影响线求解", "每个横向车道工况单独求纵向效应，最大和最小结果再形成包络。", "计算书第4章，表4-2；表4-3", "说明车道中心、横向倍率和包络过程。")
    cases = traffic["transverse_cases"]
    _table(slide, "表4-2 横向车道工况", ["工况", "启用车道", "车道中心 / m", "竖向倍率"], [[v["name"], "、".join(map(str, v["active_lanes"])), "、".join(f"{x:.2f}" for x in v["lane_centres_m"]), f"{v['vertical_multiplier']:.3f}"] for v in cases], h=3.85, widths=[3.25, 1.85, 4.2, 2.79], font_size=11)
    _textbox(slide, 0.75, 5.95, 11.7, 0.72, "▍流程  单位荷载影响线 → 车道工况 → 纵向最不利布载 → 工况间包络。\n▍用途  同一汽车包络同时进入弯矩、剪力、反力和位移组合。", 13, DARK, False, "SimSun")

    slide = new("汽车荷载内力包络", "汽车荷载弯矩和剪力的正负极值对应不同布载位置，必须分别保留。", "计算书第4章，表4-3", "读取汽车弯矩与剪力控制值，强调极值不可直接相加。")
    _xy_chart(slide, "图4-1 单幅汽车荷载弯矩包络（未计冲击）", [("最大值", a["x"], np.asarray(a["M_live_max"]) / 1000), ("最小值", a["x"], np.asarray(a["M_live_min"]) / 1000)], "弯矩 / (MN·m)")

    slide = new("作用组合与系数", "基本、标准、频遇和准永久组合分别服务于承载能力与正常使用验算。", "计算书第5章，表5-1", "逐项解释组合用途，避免把不同极限状态混用。")
    _table(slide, "表5-1 作用组合", ["组合", "永久作用", "汽车作用", "冲击", "主要用途"], [[v["name"], "/".join(f"{x:g}" for x in v["gamma_g_options"]), f"{v['vehicle_factor']*v['gamma_q']:g}", "计入" if v["include_impact"] else "不计", "承载能力" if v["id"] == "ULS_BASIC" else "正常使用"] for v in result["combinations"]], h=3.95, widths=[3.1, 1.8, 1.7, 1.3, 4.19], font_size=11)
    _textbox(slide, 0.75, 6.02, 11.7, 0.56, "Mᵤ = γ₀·max{γG·MG + γQ·MQ}；标准、频遇与准永久组合按相应汽车作用系数分别计算。", 13, INDIGO, True, "SimSun")

    slide = new("基本组合内力包络", "控制负弯矩位于主墩附近，控制正弯矩位于跨内，二者对应不同验算截面。", "计算书第5章，图5-1；表5-4、表5-6", "先读控制值，再说明正负弯矩与剪力包络的关系。")
    _xy_chart(slide, "图5-1 单幅基本组合弯矩包络", [("最大值", a["x"], np.asarray(a["M_uls_max"]) / 1000), ("最小值", a["x"], np.asarray(a["M_uls_min"]) / 1000), ("恒载", a["x"], np.asarray(a["M_dead"]) / 1000)], "弯矩 / (MN·m)")

    crit = summary["critical_effects"]
    slide = new("控制截面与支反力", "P1/P2 主墩负弯矩和剪力控制纵向验算，跨内正弯矩控制下缘受拉方向。", "计算书第5章，表5-2、表5-7、表5-8", "指出控制位置并区分反力包络极值。")
    _table(slide, "表5-2 主要控制作用效应", ["控制量", "数值", "位置"], [["基本组合负弯矩", f"{crit['negative_moment_knm']/1000:.3f} MN·m", f"x={crit['negative_moment_x_m']:.2f} m"], ["基本组合正弯矩", f"{crit['positive_moment_knm']/1000:.3f} MN·m", f"x={crit['positive_moment_x_m']:.2f} m"], ["基本组合控制剪力", f"{crit['shear_kn']/1000:.3f} MN", f"x={crit['shear_x_m']:.2f} m"], ["恒载反力合计", f"{sum(rx['dead']):.3f} kN", "A、P1、P2、A"], ["恒载平衡误差", f"{result['verification']['force_balance_error_kn']:.3e} kN", "全桥"]], h=4.1, widths=[3.3, 4.0, 4.79], font_size=11.5)

    slide = new("预应力材料与钢束初估", "钢束数量按频遇组合需求取整并形成平滑包络，截面有效束数随纵向位置变化。", "计算书第6章，表6-1", "说明这是全桥纵向等效钢束组，不替代局部锚固构造。")
    _table(slide, "表6-1 预应力材料及初估配束", ["项目", "计算采用值"], [["钢绞线", f"Φ{p['strand_diameter_mm']:.1f}；{p['strand_area_mm2']:.1f} mm²/根；fpk={p['strand_fpk_mpa']:.0f} MPa"], ["弹性模量 / 张拉控制应力", f"{p['strand_E_mpa']:.0f} / {p['jacking_stress_mpa']:.2f} MPa"], ["每束根数 / 束数 / 总根数", f"{p['strands_per_tendon']} / {p['tendon_count']} / {p['strand_count']}"], ["总钢束面积 / 张拉力", f"{p['total_area_mm2']:.0f} mm² / {p['jacking_force_kn']:.2f} kN"], ["最大有效力需求", f"{p['required_effective_force_max_kn']:.2f} kN"]], h=4.05, widths=[4.5, 7.59], font_size=11.5)
    _textbox(slide, 0.75, 6.1, 11.7, 0.52, "▍取整  钢束以整束数量布置，并保持对称跨径下的纵向包络对称。", 13, DARK, False, "SimSun")
    end_rule = p.get("end_anchor_rule")
    if end_rule:
        mode = "自动截面核心范围" if end_rule["mode"] == "auto_kern" else "用户显式输入"
        end_rows = "；".join(
            f"x={x:.3f} m：e={e:.4f} m，核心范围[{lower:.4f}, {upper:.4f}] m"
            for x, e, lower, upper in zip(end_rule["x_m"], end_rule["eccentricity_m"],
                                          end_rule["kern_lower_m"], end_rule["kern_upper_m"]))
        slide.notes_slide.notes_text_frame.text += (
            f"\n端部配束来源：prestress.end_anchor_rule；模式={end_rule['mode']}（{mode}）。"
            f"实际端部值：{end_rows}。边界按实际端截面 -W_b/A ≤ e_end ≤ W_t/A 计算；"
            "应力验算限值保持原值，显式输入按真实结果判定，不替代局部锚固区设计。")

    profile = p["profile"]; loss = p["loss_summary"]
    slide = new("预应力损失与有效应力", "摩阻、锚缩、弹性压缩、收缩徐变和松弛共同决定长期有效预应力。", "计算书第6章，表6-2、表6-3，图6-1", "指出平均总损失率、长期有效力范围和应力曲线。")
    _xy_chart(slide, "图6-1 单幅钢束应力沿程变化", [("张拉控制", profile["x"], profile["sigma_jack"]), ("摩阻后", profile["x"], profile["sigma_after_friction"]), ("长期有效", profile["x"], profile["sigma_effective"])], "钢束应力 / MPa", x=0.72, y=1.88, w=8.0, h=4.65)
    _textbox(slide, 9.02, 1.95, 3.45, 3.95, f"▍摩阻  {loss['friction_mpa_mean']:.2f} MPa\n▍锚缩  {loss['anchorage_mpa_mean']:.2f} MPa\n▍弹性压缩  {loss['elastic_shortening_mpa_mean']:.2f} MPa\n▍总损失  {loss['total_mpa_mean']:.2f} MPa\n▍平均损失率  {loss['total_ratio_mean']*100:.2f}%\n▍长期有效力  {p['effective_force_min_kn']:.0f}～{p['effective_force_max_kn']:.0f} kN", 12.5, DARK, False, "SimSun")

    slide = new("双主墩对称悬臂施工模型", "V0.12 分别计算 P1 与 P2 施工序列，并按每阶段绝对最不利结果确定控制主墩。", "计算书第7章，7.1节；表7-1", "说明双主墩分别建模与逐阶段控制逻辑，说明双主墩逐阶段计算。")
    _table(slide, "表7-1 双主墩施工模型参数", ["项目", "采用值"], [["计算主墩", "P1、P2 分别计算"], ["标准节段长度", f"{cs['segment_length_m']:.2f} m"], ["每端挂篮等效荷载", f"{cs['form_traveler_load_kn']:.2f} kN"], ["悬臂节段阶段数", f"{cs['segment_stage_count']}"], ["附加阶段", "合龙体系转换、成桥标准组合"], ["控制规则", "每阶段比较 P1/P2 上下缘应力绝对值"]], h=4.15, widths=[4.0, 8.09], font_size=11.5)
    _textbox(slide, 0.75, 6.14, 11.7, 0.46, "Mnet = MG,c − Pe；截面应力按根部实际截面特性计算。", 13, INDIGO, True, "SimSun")

    governing = cs["stages"]; stage_idx = list(range(len(governing)))
    slide = new("施工阶段内力与截面应力", f"双主墩逐阶段比较后，短暂状况混凝土压应力验算判定为 {ck['temporary_concrete_stress']['status']}。", "计算书第7章，表7-1、表7-2，图7-1", "说明控制主墩、控制阶段、最大应力和限值。")
    _xy_chart(slide, "图7-1 双主墩逐阶段控制截面应力", [("上缘", stage_idx, [v["top_stress_mpa"] for v in governing]), ("下缘", stage_idx, [v["bottom_stress_mpa"] for v in governing])], "截面应力 / MPa", x=0.72, y=1.88, w=8.0, h=4.65, x_title="施工阶段序号")
    _textbox(slide, 9.02, 2.02, 3.4, 3.7, f"▍控制主墩  {cs['governing_pier_id']}\n▍控制阶段  {cs['governing_stage_name']}\n▍最大应力绝对值  {cs['max_abs_stage_stress_mpa']:.2f} MPa\n▍压应力限值  {ck['temporary_concrete_stress']['compression_limit_mpa']:.2f} MPa\n▍判定  {ck['temporary_concrete_stress']['status']}", 13, DARK, False, "SimSun")

    rb = sv["reinforcement"]; shear = ck["shear"]
    slide = new("普通钢筋与箍筋初配", "普通钢筋从输入配筋率起步并按抗弯需求迭代，箍筋按抗剪需求取整配置。", "计算书第8章，表8-1、表8-2", "说明钢筋结果是课程设计初配并指出全桥统一配置假设。")
    _table(slide, "表8-2 普通钢筋与箍筋初配", ["部位", "初配结果", "计算说明"], [["上缘纵筋", f"{rb['top_bar_count']}Φ{rb['bar_diameter_mm']:.0f}；As={rb['top_area_mm2']:.0f} mm²", "负弯矩方向"], ["下缘纵筋", f"{rb['bottom_bar_count']}Φ{rb['bar_diameter_mm']:.0f}；As={rb['bottom_area_mm2']:.0f} mm²", "正弯矩方向"], ["双腹板箍筋", f"{shear['stirrup']['total_legs']}Φ{shear['stirrup']['diameter_mm']:.0f}@{shear['stirrup']['spacing_mm']:.0f}", "按抗剪需求取整"], ["钢筋等级", c["rebar_grade"], f"fsd={rb['fsd_mpa']:.0f} MPa"]], h=3.8, widths=[2.6, 5.2, 4.29], font_size=11.5)
    _textbox(slide, 0.75, 5.83, 11.7, 0.7, "▍假设  上、下缘纵筋沿全桥统一配置；实际工程仍需分区、构造和锚固设计。", 13, DARK, False, "SimSun")

    flex = ck["flexure"]
    slide = new("承载能力极限状态验算", f"正负弯矩抗弯判定 {flex['status']}，斜截面抗剪判定 {shear['status']}，控制结果见最大利用率。", "计算书第8章，表8-3，图8-1", "展示需求、承载力、利用率和判定。")
    categories = ["正弯矩抗弯", "负弯矩抗弯", "斜截面抗剪", "腹板尺寸上限"]
    values = [flex["positive"]["max_utilization"], flex["negative"]["max_utilization"], shear["max_utilization"], shear["web_dimension_max_utilization"]]
    _bar_chart(slide, "图8-1 承载能力验算利用率", categories, [("利用率", values), ("限值", [1, 1, 1, 1])], "利用率", x=0.72, y=1.9, w=7.8, h=4.55)
    _textbox(slide, 8.82, 2.0, 3.65, 3.9, f"▍正弯矩  {flex['positive']['demand_knm']/1000:.2f} / {flex['positive']['capacity_knm']/1000:.2f} MN·m\n▍负弯矩  {flex['negative']['demand_knm']/1000:.2f} / {flex['negative']['capacity_knm']/1000:.2f} MN·m\n▍抗剪  {shear['demand_kn']:.1f} / {shear['capacity_kn']:.1f} kN\n▍判定  {flex['status']} / {shear['status']}", 12.5, DARK, False, "SimSun")

    crack = ck["crack_resistance"]; conc = ck["concrete_stress"]; defl = ck["deflection"]
    slide = new("正常使用极限状态验算", f"抗裂 {crack['status']}，使用阶段压应力 {conc['status']}，活载挠度 {defl['status']}；各跨分别验算。", "计算书第9章，表9-1、表9-2、表9-3", "说明最大拉应力、压应力、挠度和建议预拱度。")
    _table(slide, "表9-1～表9-3 正常使用验算摘要", ["验算项目", "最大值", "限值", "判定"], [["正截面边缘拉应力", f"{crack['maximum_edge_tension_mpa']:.3f} MPa", f"{crack['edge_limit_mpa']:.3f} MPa", crack["edge_status"]], ["腹板主拉应力", f"{crack['maximum_principal_tension_mpa']:.3f} MPa", f"{crack['principal_limit_mpa']:.3f} MPa", crack["principal_status"]], ["使用阶段压应力", f"{conc['maximum_compression_mpa']:.3f} MPa", f"{conc['compression_limit_mpa']:.3f} MPa", conc["status"]], ["活载挠度控制", f"{max(v['utilization'] for v in defl['span_checks']):.3f} 利用率", "1.000", defl["status"]], ["施工阶段压应力", f"{ck['temporary_concrete_stress']['maximum_compression_mpa']:.3f} MPa", f"{ck['temporary_concrete_stress']['compression_limit_mpa']:.3f} MPa", ck['temporary_concrete_stress']["status"]]], h=4.18, widths=[3.3, 3.0, 2.8, 2.99], font_size=11)
    _textbox(slide, 0.75, 6.2, 11.7, 0.42, "▍位移  标准组合总位移已叠加长期有效预应力反拱。", 13, DARK, False, "SimSun")

    slide = new("验算汇总、追溯与敏感性", "六项具名验算全部可追溯至输入、公式、中间值和结果，参数扰动另行复算。", "计算书第9章表9-4；第10章表10-1、表10-2", "讲清可追溯记录和恒载/汽车荷载敏感性。")
    sens = result["sensitivity"]
    _bar_chart(slide, "图10-1 参数敏感性复算", [v["name"] for v in sens], [("弯矩变化率", [v["moment_change_ratio"] * 100 for v in sens]), ("剪力变化率", [v["shear_change_ratio"] * 100 for v in sens])], "相对基准变化 / %", x=0.72, y=1.9, w=7.75, h=4.55)
    _textbox(slide, 8.78, 1.95, 3.7, 4.2, f"▍具名验算  {sv['summary']['passed']}/{sv['summary']['total']} PASS\n▍计算记录  {len(result['calculation_records'])} 条\n▍追溯索引  {len(result['calculation_trace'])} 条\n▍外部成果  JSON + CSV + DOCX + PDF + DWG + PPTX\n▍数据原则  图表、表格和结论均取自同一设计状态", 12.5, DARK, False, "SimSun")

    slide = new("结论与局限", "本方案完成课程范围内的纵向计算与出图，但不替代空间、局部和下部结构工程设计。", "计算书第9.4、9.5节及第10章", "用三句话收束：做了什么、结论是什么、哪些内容不在本次范围。")
    _textbox(slide, 0.82, 1.75, 11.7, 3.95, f"▍方案  {g['left_side_span']:g}+{g['spans'][1]:g}+{g['right_side_span']:g} m {desc['title']}，{desc['arrangement']}。\n▍计算  恒载、汽车荷载、组合包络、预应力损失与主次效应、P1/P2 双主墩施工阶段、钢筋初配、承载能力和正常使用验算。\n▍结果  六项具名验算 {sv['summary']['passed']}/{sv['summary']['total']} 通过；课程状态为 {result['status_code']}。\n▍局限  横向板壳、空间扭转、锚固区局部实体、风、地震、下部结构及基础承载力属于课程范围外。", 16, DARK, False, "SimSun")
    cc=result.get('crosscheck')
    if cc:
        _textbox(slide, .82, 5.9, 11.7, .42, f"独立复核 {cc['status']}｜数值 {cc['numerical_status']}｜固定方案 {cc['design_status']}｜覆盖 PARTIAL", 12, RED, True, 'SimHei')
        slide.notes_slide.notes_text_frame.text+='\n独立复核来源：计算书第11章；复核报告JSON；run_id='+cc['provenance']['run_id']+'。'+cc['message']
    _textbox(slide, 1.6, 6.45, 10.1, 0.38, "所有结论均可回溯至计算书表号、设计结果 JSON 与计算追溯 CSV。", 12, RED, True, "SimHei", align=PP_ALIGN.CENTER)

    assert len(prs.slides) == SLIDE_COUNT
    prs.core_properties.title = f"{c['project_name']}课程设计汇报"
    prs.core_properties.subject = "23页学术答辩汇报；数据来自BridgeGenerator V0.12计算结果"
    prs.core_properties.author = "BridgeGenerator V0.12"
    prs.save(path)
    _finalize_chart_package(path)

